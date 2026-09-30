"""Profile the concepts of a material: catalog, extract, check, record."""

import argparse
import asyncio
from contextlib import asynccontextmanager
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import time

from mcp import ClientSession
from mcp import StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import load_workbook
import yaml

from wiki_consistency import evidence
from wiki_consistency.evidence import _payload_path
from wiki_consistency.instance import instance_path
from wiki_consistency.instance import revisions
from wiki_consistency.instance import roots

LIMIT = 200 * 1024**2
BACKFIRE = Path(__file__).resolve().parents[5] / "packages/backfire"
TICKED = re.compile(r'^- \[[xX]\] (\d+) ("(?:[^"\\]|\\.)*")', re.M)
OUTCOMES = {
    ("auto", "verified"): "kept",
    ("auto", "contradicted"): "dropped",
    ("auto", "unsupported"): "dropped",
}


def _run_dir(name):
    state = os.environ.get("XDG_STATE_HOME") or Path.home() / ".local/state"
    return Path(state) / "verbose-broccoli/concept-profile" / name


def _budget(run, text):
    used = sum(p.stat().st_size for p in run.rglob("*") if p.is_file())
    if used + len(text.encode()) > LIMIT:
        raise ValueError(f"run folder budget of {LIMIT} bytes exceeded")


def _write(path, text, run=None):
    if run is not None:
        _budget(run, text)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(text, encoding="utf-8")
    temporary.replace(path)


def _jsonl(path):
    return [json.loads(line) for line in path.read_text().splitlines()]


def _norm(text):
    return " ".join(text.split())


def _catalog(root, page):
    """Return the catalog page's metadata and its concepts by key."""
    text = (root / "wiki" / page).read_text(encoding="utf-8")
    metadata = yaml.safe_load(text.split("---\n", 2)[1])
    spec, source = metadata["catalog"], metadata["sources"][0]
    item = next(
        item
        for item in revisions(root)[source["id"]]
        if item["revision"] == source["revision"]
    )
    book = load_workbook(_payload_path(root, item), read_only=True)
    rows = book[spec["sheet"]].iter_rows(values_only=True)
    header, columns = list(next(rows)), spec["columns"]

    def cell(row, name):
        return str(row[header.index(name)] or "").strip()

    entries = {
        key: {
            "id": cell(row, columns["id"]),
            "level": cell(row, columns["level"]),
            "label": " / ".join(cell(row, n) for n in columns["label"]),
            "statement": cell(row, columns["statement"]),
            "examples": cell(row, columns["examples"]),
        }
        for key, row in enumerate(rows, 1)
    }
    return metadata, entries


def _catalog_command(args, root, run):
    entries = _catalog(root, args.catalog)[1]
    fields = ("level", "label", "statement")
    lines = (
        "\t".join([str(key)] + [_norm(e[name]) for name in fields])
        for key, e in entries.items()
    )
    _write(run / "catalog.tsv", "\n".join(lines) + "\n", run)


def _extract_command(args, root, run):
    cache = roots(os.environ)["cache"]
    for source in args.source:
        item = revisions(root)[source][-1]
        payload = _payload_path(root, item)
        if payload.suffix.lower() == ".pdf":
            text = subprocess.check_output(
                ["pdftotext", "-raw", payload, "-"], text=True
            )
        else:
            evidence.convert(root, args.wiki, cache, {source: [item]})
            found = evidence.read(cache, args.wiki, source, item["revision"])
            text = found.get("text", "")
        if not any(c.isalpha() for c in text):
            raise ValueError(f"{source}: no text; a scan needs the user")
        _write(run / f"text/{source}.txt", text, run)


@asynccontextmanager
async def _backfire():
    server = StdioServerParameters(
        command="uv",
        args=["--directory", str(BACKFIRE), "run", "--frozen", "--offline"]
        + ["--no-sync", "backfire", "serve-mcp", "--education"],
    )
    async with stdio_client(server) as (read, write):
        async with ClientSession(read, write) as session:
            await session.initialize()
            yield session


async def _verify(session, arguments):
    result = await session.call_tool("jev_verify", arguments)
    if result.isError:
        raise ValueError(result.content[0].text)
    return json.loads(result.content[0].text)


def _refusals(run, proposals, entries):
    texts = {
        p.stem: _norm(p.read_text(encoding="utf-8"))
        for p in (run / "text").glob("*.txt")
    }
    for n, row in enumerate(proposals, 1):
        if _norm(row["text"]) not in texts.get(row["source"], ""):
            yield f"proposal {n}: sentence is absent from the extracted text"
        elif not set(row["concepts"]) <= entries.keys():
            yield f"proposal {n}: concept key is not in the catalog"


async def _send(run, todo, entries, claim, done):
    async with _backfire() as session:
        for n, row in todo:
            started = time.monotonic()
            found = [entries[key] for key in sorted(set(row["concepts"]))]
            text = _norm(row["text"])
            parts = ("label", "statement", "examples")
            items = [{"id": "sentence", "text": text}] + [
                {"id": e["id"], "text": "\n".join(e[p] for p in parts)}
                for e in found
            ]
            claims = [claim.format(text=text, **e) for e in found]
            arguments = {"claims": claims, "evidence": items}
            response = await _verify(session, arguments) if found else {}
            results = response.get("results", [])
            if len(results) != len(found):
                raise ValueError(
                    f"{len(results)} results for {len(found)} claims"
                )
            check = {
                "row": n,
                **{k: response.get(k) for k in ("provider", "model")},
                "usage": response.get("usage") or {},
                "seconds": time.monotonic() - started,
                "results": [
                    {
                        **r,
                        "concept": e["id"],
                        "outcome": OUTCOMES.get(
                            (r.get("action"), r.get("verdict")), "unclear"
                        ),
                    }
                    for e, r in zip(found, results)
                ],
            }
            line = json.dumps(check) + "\n"
            _budget(run, line)
            with (run / "checks.jsonl").open("a", encoding="utf-8") as stream:
                stream.write(line)
            done.append(check)


def _check_command(args, root, run):
    metadata, entries = _catalog(root, args.catalog)
    proposals = _jsonl(run / "proposals.jsonl")
    if refused := list(_refusals(run, proposals, entries)):
        print(*refused, sep="\n", file=sys.stderr)
        return 1
    path = run / "checks.jsonl"
    lines = path.read_bytes().split(b"\n") if path.exists() else [b""]
    if lines[-1]:  # A torn last line has no newline: drop it.
        path.write_bytes(b"".join(line + b"\n" for line in lines[:-1]))
    done = [json.loads(line) for line in lines[:-1]]
    todo = [(n, r) for n, r in enumerate(proposals, 1) if n > len(done)]
    if todo:
        claim = metadata["catalog"]["claim"]
        try:
            asyncio.run(_send(run, todo, entries, claim, done))
        except Exception as error:  # noqa: BLE001 - any backfire failure.
            print(f"row {len(done) + 1}: {error}", file=sys.stderr)
            return 2
    by_id = {e["id"]: e for e in entries.values()}
    sheet = []
    for check in done:
        n = check["row"]
        unclear = [
            r["concept"] for r in check["results"] if r["outcome"] == "unclear"
        ]
        if unclear:
            sheet.append(f"## {n}. {_norm(proposals[n - 1]['text'])}\n")
            for i in unclear:
                label, statement = by_id[i]["label"], by_id[i]["statement"]
                sheet.append(
                    f"- [ ] {n} {json.dumps(i)} — {label}: {statement}"
                )
            sheet.append("")
    _write(run / "review.md", "\n".join(sheet), run)
    results = [r for c in done for r in c["results"]]
    report = {
        "sentences": len(proposals),
        "calls": sum(bool(c["results"]) for c in done),
        "invalid": sum(r.get("verdict") in (None, "unknown") for r in results),
        "tokens": sum(
            c["usage"].get("input_tokens", 0)
            + c["usage"].get("output_tokens", 0)
            for c in done
        ),
        "seconds": round(sum(c["seconds"] for c in done), 1),
    }
    for outcome in ("kept", "dropped", "unclear"):
        report[outcome] = sum(r["outcome"] == outcome for r in results)
    print(json.dumps(report))
    return int(report["invalid"] > 0)


def _record_command(args, root, run):
    metadata, entries = _catalog(root, args.catalog)
    proposals = _jsonl(run / "proposals.jsonl")
    checks = _jsonl(run / "checks.jsonl")
    if [c["row"] for c in checks] != list(range(1, len(proposals) + 1)):
        raise LookupError("checks.jsonl does not cover every proposal")
    order = {e["id"]: key for key, e in entries.items()}
    sheet = run / "review.md"
    sheet = sheet.read_text(encoding="utf-8") if args.reviewed else ""
    ticked = {(int(n), json.loads(i)) for n, i in TICKED.findall(sheet)}
    rows = []
    for n, (proposal, check) in enumerate(zip(proposals, checks), 1):
        by = {
            o: [r["concept"] for r in check["results"] if r["outcome"] == o]
            for o in ("kept", "unclear")
        }
        accepted = [i for i in by["unclear"] if (n, i) in ticked]
        rows.append(
            {
                "n": n,
                "source": proposal["source"],
                "part": proposal["part"],
                "text": _norm(proposal["text"]),
                "concepts": sorted(by["kept"] + accepted, key=order.get),
                "unclear": [] if args.reviewed else by["unclear"],
            }
        )
    kept = sum(len(r["concepts"]) for r in rows)
    unclear = sum(len(r["unclear"]) for r in rows)
    proposed = sum(len(c["results"]) for c in checks)
    sources = [
        {"id": s, "revision": revisions(root)[s][-1]["revision"]}
        for s in dict.fromkeys(p["source"] for p in proposals)
    ]
    checker = {f"{c['provider']} {c['model']}" for c in checks if c["provider"]}
    page = {
        "title": args.title,
        "summary": args.summary,
        "topics": metadata["topics"],
        "sources": sources + metadata["sources"],
        "profile": {
            "catalog": f"../{args.catalog}",
            "data": f"{args.name}.jsonl",
            "proposer": args.proposer,
            "checker": "; ".join(sorted(checker)),
            "counts": {
                "sentences": len(rows),
                "kept": kept,
                "dropped": proposed - kept - unclear,
                "unclear": unclear,
            },
        },
    }
    body = (
        f"{args.title}: every sentence and its concepts, from the "
        f"[catalog](../{args.catalog}); rows in "
        f"[{args.name}.jsonl]({args.name}.jsonl).\n"
    )
    front = yaml.safe_dump(page, sort_keys=False, allow_unicode=True)
    data = "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows)
    _write(root / f"wiki/profiles/{args.name}.jsonl", data)
    _write(root / f"wiki/profiles/{args.name}.md", f"---\n{front}---\n\n{body}")


def main(argv=None):
    """Run one command and return its exit code."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--wiki", default="work")
    parser.add_argument("--run", required=True)
    commands = parser.add_subparsers(dest="command", required=True)

    def command(name, function, *options):
        sub = commands.add_parser(name)
        sub.set_defaults(function=function)
        for option in options:
            sub.add_argument(f"--{option}", required=True)
        return sub

    command("catalog", _catalog_command, "catalog")
    command("check", _check_command, "catalog")
    names = ("catalog", "name", "title", "summary", "proposer")
    record = command("record", _record_command, *names)
    record.add_argument("--reviewed", action="store_true")
    extract = command("extract", _extract_command)
    extract.add_argument("--source", nargs="+", required=True)
    args = parser.parse_args(argv)
    try:
        root = instance_path(args.wiki, os.environ)
        return args.function(args, root, _run_dir(args.run)) or 0
    except (OSError, ValueError, LookupError, subprocess.SubprocessError) as e:
        print(f"concept-profile: {e}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
