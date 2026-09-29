from importlib.metadata import version
import io
import json
import os
import signal
import socket
import warnings
import zipfile

from hwpx import Hwp5ConversionWarning
from hwpx import HwpxDocument
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
import pytest

from wiki_consistency import evidence


def _docx_bytes(path):
    content_types = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">\n'
        b'  <Default Extension="rels" ContentType="'
        b"application/vnd.openxmlformats-package."
        b'relationships+xml"/>\n'
        b'  <Default Extension="xml" ContentType="application/xml"/>\n'
        b'  <Override PartName="/word/document.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.'
        b'wordprocessingml.document.main+xml"/>\n'
        b"</Types>"
    )
    relationships = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        b'<Relationships xmlns="'
        b'http://schemas.openxmlformats.org/package/2006/relationships">\n'
        b'  <Relationship Id="rId1" Type="'
        b"http://schemas.openxmlformats.org/officeDocument/2006/"
        b'relationships/officeDocument" Target="word/document.xml"/>\n'
        b"</Relationships>"
    )
    document = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        b'<w:document xmlns:w="'
        b'http://schemas.openxmlformats.org/wordprocessingml/2006/main">\n'
        b"  <w:body><w:p><w:r><w:t>Synthetic DOCX evidence</w:t></w:r>"
        b"</w:p><w:sectPr/></w:body>\n"
        b"</w:document>"
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document)


def _pptx_bytes(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    shape.text_frame.text = "Synthetic PPTX evidence"
    presentation.save(path)


def _pdf_bytes(stream):
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, value in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode())
        output.extend(value + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Root 1 0 R /Size {len(offsets)} >>\n"
        f"startxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def _revision(instance, source_id, revision, filename, payload):
    bag = instance / "raw" / "files" / source_id / revision
    payload_path = bag / "data" / filename
    payload_path.parent.mkdir(parents=True, exist_ok=True)
    payload_path.write_bytes(payload)
    return {
        "kind": "files",
        "id": source_id,
        "revision": revision,
        "path": f"raw/files/{source_id}/{revision}",
    }


def _revisions(*items):
    return {item["id"]: [item] for item in items}


def _evidence_root(cache, wiki_id):
    return (
        cache
        / "wiki-evidence"
        / wiki_id
        / f"markitdown-{evidence.CONVERTER_VERSION}"
    )


def _chatgpt_conversation(number=1):
    user_id = f"synthetic-user-{number}"
    assistant_id = f"synthetic-assistant-{number}"
    return {
        "title": f"Synthetic conversation {number}",
        "mapping": {
            user_id: {
                "id": user_id,
                "message": {
                    "id": user_id,
                    "author": {"role": "user"},
                    "content": {
                        "content_type": "text",
                        "parts": [f"합성 질문 {number}: Where is a library?"],
                    },
                },
                "parent": None,
                "children": [assistant_id],
            },
            assistant_id: {
                "id": assistant_id,
                "message": {
                    "id": assistant_id,
                    "author": {"role": "assistant"},
                    "content": {
                        "content_type": "text",
                        "parts": [
                            f"합성 답변 {number}: It is a place to read."
                        ],
                    },
                },
                "parent": user_id,
                "children": [],
            },
        },
        "current_node": assistant_id,
    }


def _chatgpt_export_bytes(*, ensure_ascii=True, numbered=False):
    conversations = [
        _chatgpt_conversation(number)
        for number in ([1, 2] if numbered else [1])
    ]
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        if numbered:
            for number, conversation in enumerate(conversations, start=1):
                archive.writestr(
                    f"conversation_{number}.json",
                    json.dumps(conversation, ensure_ascii=ensure_ascii),
                )
        else:
            archive.writestr(
                "conversations.json",
                json.dumps(conversations, ensure_ascii=ensure_ascii),
            )
        archive.writestr(
            "chat.html", "<html><body>Synthetic export</body></html>"
        )
        archive.writestr("user.json", '{"display_name":"Synthetic user"}')
    return output.getvalue()


def _no_socket(*args, **kwargs):
    del args, kwargs  # Unused.
    raise AssertionError("evidence conversion must not open a socket")


def test_convert_supported_files_and_pass_through(tmp_path, monkeypatch):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    docx = tmp_path / "source.docx"
    pptx = tmp_path / "source.pptx"
    _docx_bytes(docx)
    _pptx_bytes(pptx)

    items = [
        _revision(instance, "docx", "r1", docx.name, docx.read_bytes()),
        _revision(instance, "pptx", "r1", pptx.name, pptx.read_bytes()),
        _revision(
            instance,
            "pdf",
            "r1",
            "source.pdf",
            _pdf_bytes(b"BT /F1 12 Tf 72 72 Td (Synthetic PDF evidence) Tj ET"),
        ),
        _revision(instance, "text", "r1", "source.txt", b"Plain text evidence"),
        _revision(
            instance, "markdown", "r1", "source.md", b"# Markdown evidence\n"
        ),
    ]
    monkeypatch.setattr(socket, "socket", _no_socket)

    result = evidence.convert(instance, "wiki-a", cache, _revisions(*items))

    assert result["converted"] == 5
    assert result["present"] == 0
    assert result["unreadable"] == []
    assert (
        "Synthetic DOCX evidence"
        in evidence.read(cache, "wiki-a", "docx", "r1")["text"]
    )
    assert (
        "Synthetic PPTX evidence"
        in evidence.read(cache, "wiki-a", "pptx", "r1")["text"]
    )
    assert (
        "Synthetic PDF evidence"
        in evidence.read(cache, "wiki-a", "pdf", "r1")["text"]
    )
    assert (
        evidence.read(cache, "wiki-a", "text", "r1")["text"]
        == "Plain text evidence"
    )
    assert (
        evidence.read(cache, "wiki-a", "markdown", "r1")["text"]
        == "# Markdown evidence\n"
    )
    assert {path.relative_to(cache).parts[0] for path in cache.rglob("*")} == {
        "wiki-evidence"
    }


def test_convert_xlsx_hwp_and_hwpx_sources(tmp_path, monkeypatch):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()

    xlsx = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Synthetic XLSX evidence"
    workbook.save(xlsx)
    workbook.close()

    hwp = tmp_path / "source.hwp"
    hwpx = tmp_path / "source.hwpx"
    for path in (hwp, hwpx):
        document = HwpxDocument.new()
        document.add_paragraph(f"Synthetic {path.suffix[1:].upper()} evidence")
        document.save_to_path(path)
        document.close()

    items = [
        _revision(instance, "xlsx", "r1", xlsx.name, xlsx.read_bytes()),
        _revision(instance, "hwp", "r1", hwp.name, hwp.read_bytes()),
        _revision(instance, "hwpx", "r1", hwpx.name, hwpx.read_bytes()),
    ]
    monkeypatch.setattr(socket, "socket", _no_socket)

    result = evidence.convert(instance, "wiki-a", cache, _revisions(*items))

    assert result["unreadable"] == []
    for source_id in ("xlsx", "hwp", "hwpx"):
        assert (
            f"Synthetic {source_id.upper()} evidence"
            in evidence.read(cache, "wiki-a", source_id, "r1")["text"]
        )


def test_unreadable_revisions_record_the_reason(tmp_path):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    damaged = tmp_path / "damaged.docx"
    with zipfile.ZipFile(damaged, "w") as archive:
        archive.writestr("word/document.xml", "<broken>")
    items = [
        _revision(
            instance,
            "unsupported",
            "r1",
            "unsupported.foo",
            b"\xff\xfe\x00\x00",
        ),
        _revision(
            instance,
            "hwp",
            "r1",
            "damaged.hwp",
            b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",
        ),
        _revision(
            instance,
            "image-pdf",
            "r1",
            "image.pdf",
            _pdf_bytes(b"q 0 0 1 rg 72 72 100 100 re f Q"),
        ),
        _revision(
            instance, "damaged", "r1", "damaged.docx", damaged.read_bytes()
        ),
    ]

    result = evidence.convert(instance, "wiki-a", cache, _revisions(*items))

    reasons = {entry["id"]: entry["reason"] for entry in result["unreadable"]}
    assert reasons == {
        "unsupported": "unsupported_format",
        "hwp": "conversion_failed",
        "image-pdf": "empty_text",
        "damaged": "conversion_failed",
    }
    for source_id, reason in reasons.items():
        mark = (
            _evidence_root(cache, "wiki-a") / source_id / "r1.unreadable.json"
        )
        stored = json.loads(mark.read_text())
        assert stored["reason"] == reason
        assert isinstance(stored["detail"], str)
        assert evidence.read(cache, "wiki-a", source_id, "r1") == {
            "unreadable": reason
        }


def test_hwp5_conversion_warning_is_marked_and_reused(tmp_path, monkeypatch):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()

    document = HwpxDocument.new()
    document.add_paragraph("Synthetic partial HWP evidence")
    hwp = tmp_path / "source.hwp"
    document.save_to_path(hwp)
    document.close()
    item = _revision(instance, "hwp", "r1.2", hwp.name, hwp.read_bytes())

    detail = "Synthetic HWP5 conversion warning"
    original_open = HwpxDocument.open.__func__

    def open_with_warning(cls, source):
        warnings.warn(detail, Hwp5ConversionWarning)
        return original_open(cls, source)

    monkeypatch.setattr(HwpxDocument, "open", classmethod(open_with_warning))

    write_immutable = evidence._write_immutable
    write_order = []

    def record_write(target, content, root, budget_bytes):
        write_order.append(target.name)
        return write_immutable(target, content, root, budget_bytes)

    monkeypatch.setattr(evidence, "_write_immutable", record_write)

    first = evidence.convert(instance, "wiki-a", cache, _revisions(item))
    partial = [{"id": "hwp", "revision": "r1.2", "detail": detail}]
    partial_mark = _evidence_root(cache, "wiki-a") / "hwp" / "r1.2.partial.json"

    assert first["partial"] == partial
    assert first["unreadable"] == []
    assert write_order == ["r1.2.partial.json", "r1.2.md"]
    assert json.loads(partial_mark.read_text()) == {"detail": detail}
    assert (
        "Synthetic partial HWP evidence"
        in evidence.read(cache, "wiki-a", "hwp", "r1.2")["text"]
    )

    second = evidence.convert(instance, "wiki-a", cache, _revisions(item))

    assert second["partial"] == partial
    assert second["present"] == 1


def test_convert_never_rewrites_and_reads_unconverted_as_missing(tmp_path):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    item = _revision(instance, "source", "r1", "source.txt", b"first content")
    revisions = _revisions(item)

    assert (
        evidence.convert(instance, "wiki-a", cache, revisions)["converted"] == 1
    )
    target = _evidence_root(cache, "wiki-a") / "source" / "r1.md"
    original = target.read_bytes()
    item_path = instance / item["path"] / "data" / "source.txt"
    item_path.write_text("changed content")

    result = evidence.convert(instance, "wiki-a", cache, revisions)

    assert result["converted"] == 0
    assert result["present"] == 1
    assert target.read_bytes() == original
    with pytest.raises(LookupError):
        evidence.read(cache, "wiki-a", "missing", "r1")


def test_budget_is_checked_before_each_write(tmp_path):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    items = [
        _revision(instance, "a", "r1", "source.txt", b"alpha"),
        _revision(instance, "b", "r1", "source.txt", b"bravo"),
    ]

    with pytest.raises(ValueError, match="wiki-evidence.*budget"):
        evidence.convert(
            instance, "wiki-a", cache, _revisions(*items), budget_bytes=7
        )

    assert (_evidence_root(cache, "wiki-a") / "a" / "r1.md").exists()
    assert not (_evidence_root(cache, "wiki-a") / "b" / "r1.md").exists()
    assert not list(cache.rglob("*.wiki-consistency-tmp"))


@pytest.mark.parametrize("signum", [signal.SIGINT, signal.SIGTERM])
def test_interrupted_write_cleans_temporary_files(
    tmp_path, monkeypatch, signum
):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    item = _revision(instance, "source", "r1", "source.txt", b"source evidence")

    def interrupt_rename(*args):
        del args  # Unused.
        os.kill(os.getpid(), signum)

    monkeypatch.setattr(evidence.os, "rename", interrupt_rename)
    with pytest.raises(SystemExit):
        evidence.convert(instance, "wiki-a", cache, _revisions(item))

    assert not list(cache.rglob("*.wiki-consistency-tmp"))
    assert not (_evidence_root(cache, "wiki-a") / "source" / "r1.md").exists()


def test_failed_rename_cleans_temporary_files(tmp_path, monkeypatch):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    item = _revision(instance, "source", "r1", "source.txt", b"source evidence")

    def fail_rename(*args):
        del args  # Unused.
        raise OSError("synthetic rename failure")

    monkeypatch.setattr(evidence.os, "rename", fail_rename)
    with pytest.raises(OSError, match="synthetic rename failure"):
        evidence.convert(instance, "wiki-a", cache, _revisions(item))

    assert not list(cache.rglob("*.wiki-consistency-tmp"))


def test_next_convert_removes_abandoned_temporary_files(tmp_path):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    stale = cache / "wiki-evidence" / ".abandoned.wiki-consistency-tmp"
    stale.parent.mkdir(parents=True)
    stale.write_text("partial")

    evidence.convert(instance, "wiki-a", cache, {})

    assert not stale.exists()


def test_chatgpt_export_unescapes_korean_and_english_messages(tmp_path):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    payload = _chatgpt_export_bytes(ensure_ascii=True)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert b"\\u" in archive.read("conversations.json")
        assert {"chat.html", "user.json", "conversations.json"} <= set(
            archive.namelist()
        )
    item = _revision(instance, "chatgpt-export", "r1", "export.zip", payload)

    result = evidence.convert(instance, "wiki-a", cache, _revisions(item))
    text = evidence.read(cache, "wiki-a", "chatgpt-export", "r1")["text"]

    assert result["unreadable"] == []
    assert "합성 질문 1: Where is a library?" in text
    assert "합성 답변 1: It is a place to read." in text


def test_chatgpt_export_with_direct_characters_keeps_messages(tmp_path):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    payload = _chatgpt_export_bytes(ensure_ascii=False)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert "합성 질문 1".encode() in archive.read("conversations.json")
    item = _revision(instance, "chatgpt-export", "r1", "export.zip", payload)

    result = evidence.convert(instance, "wiki-a", cache, _revisions(item))
    text = evidence.read(cache, "wiki-a", "chatgpt-export", "r1")["text"]

    assert result["unreadable"] == []
    assert "합성 질문 1: Where is a library?" in text
    assert "합성 답변 1: It is a place to read." in text


def test_chatgpt_export_with_numbered_conversation_json_files(tmp_path):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    payload = _chatgpt_export_bytes(ensure_ascii=True, numbered=True)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert "conversations.json" not in archive.namelist()
        assert {"conversation_1.json", "conversation_2.json"} <= set(
            archive.namelist()
        )
    item = _revision(instance, "chatgpt-export", "r1", "export.zip", payload)

    result = evidence.convert(instance, "wiki-a", cache, _revisions(item))
    text = evidence.read(cache, "wiki-a", "chatgpt-export", "r1")["text"]

    assert result["unreadable"] == []
    for number in (1, 2):
        assert f"합성 질문 {number}: Where is a library?" in text
        assert f"합성 답변 {number}: It is a place to read." in text


def test_jsonl_unescapes_records_and_keeps_blank_lines(tmp_path):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    payload = (
        b'{"text":"\\ud55c\\uad6d Alpha"}\n\n{"text":"\\uc601\\uc5b4 Bravo"}\n'
    )
    item = _revision(instance, "jsonl", "r1", "source.jsonl", payload)

    result = evidence.convert(instance, "wiki-a", cache, _revisions(item))
    text = evidence.read(cache, "wiki-a", "jsonl", "r1")["text"]

    assert result["unreadable"] == []
    assert text == '{"text": "한국 Alpha"}\n\n{"text": "영어 Bravo"}\n'


def test_invalid_json_returns_plain_text_converter_output(tmp_path):
    instance = tmp_path / "instance"
    cache = tmp_path / "cache"
    instance.mkdir()
    payload = b'{"text":"\\uac00", invalid}\n'
    item = _revision(instance, "invalid-json", "r1", "source.json", payload)
    payload_path = instance / item["path"] / "data" / "source.json"
    plain_text = evidence.MarkItDown().convert(payload_path).text_content or ""

    result = evidence.convert(instance, "wiki-a", cache, _revisions(item))

    assert result["unreadable"] == []
    assert (
        evidence.read(cache, "wiki-a", "invalid-json", "r1")["text"]
        == plain_text
    )


def test_evidence_cache_path_uses_a_local_converter_revision(tmp_path):
    path = _evidence_root(tmp_path, "wiki-a")

    assert path.name != f"markitdown-{version('markitdown')}"
    assert evidence.CONVERTER_VERSION == (
        f"{version('markitdown')}-hwpx-{version('python-hwpx')}-json-2"
    )
